"""파이프라인 추적 엑셀. python tools/trace_xlsx.py labels/auto_gold.csv labels/gold.csv ../추적.xlsx
열: image_id | 정답 | 예측 | 판정 | stage | OCR 줄(묶은 뒤) | 후보와 점수 근거 | 원시 조각(좌표)"""
import sys, os, csv, json
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from dateparse import explain
from ocr import group_lines

pred_csv, gold_csv, out = sys.argv[1:4]
gold = {r["image_id"]: r for r in csv.DictReader(open(gold_csv, encoding="utf-8"))} if os.path.exists(gold_csv) else {}
rows = list(csv.DictReader(open(pred_csv, encoding="utf-8")))
STAGE_KO = {"s1": "1단계 짧은 조각", "s2": "2단계 나머지 조각", "rot90": "회전 90", "rot270": "회전 270",
            "rot180": "회전 180", "clahe": "대비 강화", "up2x": "2배 확대", "fail": "전부 실패", "error": "예외"}
wb = Workbook(); ws = wb.active; ws.title = "추적"
hdr = ["image_id", "정답", "예측", "판정", "OCR 단계", "OCR 줄 (좌표로 묶은 뒤, 위->아래)", "후보와 선택 근거 (위가 선택)", "원시 조각 [텍스트, y, 높이, x]"]
ws.append(hdr)
for c in range(1, len(hdr) + 1):
    ws.cell(1, c).font = Font(bold=True); ws.cell(1, c).fill = PatternFill("solid", fgColor="DDDDDD")
for col, w in zip("ABCDEFGH", [10, 14, 14, 10, 14, 60, 70, 50]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
RED, YEL, GRN = [PatternFill("solid", fgColor=c) for c in ("FFC7CE", "FFEB9C", "C6EFCE")]
for i, r in enumerate(rows, start=2):
    g = gold.get(r["image_id"])
    gold_s = f'{g["year"]}-{g["month"]}-{g["day"]}' if g else ""
    pred_s = r["final_date"]
    if r.get("raw"):
        raw = json.loads(r["raw"])
        items = [(t, [[cx, cy - h / 2], [cx + 100, cy - h / 2], [cx + 100, cy + h / 2], [cx, cy + h / 2]]) for t, cy, h, cx in raw]
        lines = group_lines(items)
        raw_s = "\n".join(f"{t}  (y{cy} h{h} x{cx})" for t, cy, h, cx in raw)
    else:
        lines = [t for t in r["texts"].split(" | ") if t]; raw_s = ""
    _, why = explain(lines)
    if not g: verdict = ""
    elif gold_s == pred_s: verdict = "맞음"
    elif pred_s == "NONE": verdict = "못 찾음"
    else:
        ok = sum(a == b for a, b in zip(gold_s.split("-"), pred_s.split("-")))
        verdict = f"부분 {ok}/3"
    vals = [r["image_id"], gold_s, pred_s, verdict, STAGE_KO.get(r["stage"], r["stage"]),
            "\n".join(lines), "\n".join(why) if why else "(날짜 후보 없음)", raw_s]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(i, c, v); cell.alignment = Alignment(wrap_text=True, vertical="top")
        if c in (1, 2, 3): cell.number_format = "@"
    if verdict.startswith("맞음"): ws.cell(i, 4).fill = GRN
    elif verdict.startswith("부분"): ws.cell(i, 4).fill = YEL
    elif verdict: ws.cell(i, 4).fill = RED
wb.save(out)
print("saved", out, len(rows), "rows")
