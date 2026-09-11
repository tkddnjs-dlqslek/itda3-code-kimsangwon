# -*- coding: utf-8 -*-
"""det(검출) 파인튜닝용 박스 라벨링 스켈레톤을 만든다. 박스 라벨링 자체는 자동화 범위 밖:
PPOCRLabel(권장) 또는 LabelMe 로 사람이 직접 그려야 한다.

python finetune/runpod/prep_det_data.py [--data finetune/data]

만드는 것:
- finetune/data/det/images/<image_id>.<ext>          (hard_images.txt 이미지 복사본)
- finetune/data/det/images/Label.txt                  (PPOCRLabel 이 바로 열 수 있는 빈 라벨)
- finetune/data/det/labeling_instructions.xlsx        (라벨링 방법 + 대상 목록)

PPOCRLabel 사용법 (자동화 아님, 사람이 직접):
  pip install PPOCRLabel
  PPOCRLabel --lang ch
  -> "파일 열기" 로 finetune/data/det/images 폴더 선택 -> 각 이미지에서 소비기한 날짜가
     적힌 영역에 사각형을 그리고 텍스트를 채운 뒤 저장 -> Label.txt 가 갱신됨.
학습용 최종 포맷(PaddleOCR det, det_v5_finetune.yml 이 기대하는 것):
  <이미지경로>\t[{"transcription": "2026.06.28", "points": [[x1,y1]...[x4,y4]]}, ...]
PPOCRLabel 이 만드는 Label.txt 가 이미 이 포맷이므로, 라벨링 후
Label.txt 를 train_label.txt(또는 val_label.txt) 로 이름만 바꾸면 된다 (train/val 분할은
이미지 단위로 사람이 직접 나누거나 finetune/prep_rec_data.py 의 15% 홀드아웃 방식을 따라도 됨).
"""
from __future__ import annotations

import argparse
import json
import os
import shutil

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.dirname(os.path.dirname(_HERE))
_IMAGES_DIR = os.path.normpath(os.path.join(_REPO, os.pardir, "images"))
_EXTS = (".jpg", ".jpeg", ".png")


def _find(image_id: str) -> str | None:
    for ext in _EXTS:
        p = os.path.join(_IMAGES_DIR, image_id + ext)
        if os.path.exists(p):
            return p
    return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=os.path.join(os.path.dirname(_HERE), "data"))
    args = ap.parse_args()
    data_dir = os.path.abspath(args.data)

    hard_images_txt = os.path.join(data_dir, "rec", "hard_images.txt")
    if not os.path.exists(hard_images_txt):
        raise SystemExit(f"없음: {hard_images_txt} (prep_rec_data.py 먼저 실행)")
    image_ids = [l.strip() for l in open(hard_images_txt, encoding="utf-8") if l.strip()]

    out_images = os.path.join(data_dir, "det", "images")
    os.makedirs(out_images, exist_ok=True)

    label_lines = []
    missing = []
    for iid in image_ids:
        src = _find(iid)
        if src is None:
            missing.append(iid)
            continue
        fname = os.path.basename(src)
        dst = os.path.join(out_images, fname)
        if not os.path.exists(dst):
            shutil.copyfile(src, dst)
        label_lines.append(f"{fname}\t{json.dumps([], ensure_ascii=False)}")

    with open(os.path.join(out_images, "Label.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(label_lines) + "\n")

    _write_instructions_xlsx(data_dir, image_ids, missing)

    print(f"det 라벨링 대상: {len(label_lines)}장 (이미지 없음 {len(missing)}장)")
    print(f"이미지: {out_images}")
    print(f"안내서: {os.path.join(data_dir, 'det', 'labeling_instructions.xlsx')}")


def _write_instructions_xlsx(data_dir: str, image_ids: list, missing: list) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    ws.append(["PPOCRLabel 박스 라벨링 안내"])
    ws["A1"].font = Font(bold=True, size=14)
    steps = [
        "1) pip install PPOCRLabel",
        "2) PPOCRLabel --lang ch 실행",
        "3) 파일 > 폴더 열기 -> finetune/data/det/images 선택",
        "4) 각 이미지에서 소비기한 날짜가 적힌 영역에 사각형을 그리고, 인식된 텍스트 칸에",
        "   실제 표기(예: 2026.06.28)를 정확히 입력. 날짜가 없는 이미지는 빈 채로 저장",
        "5) 저장(Ctrl+S)하면 images/Label.txt 가 갱신됨. 이것이 그대로 det 학습 라벨",
        "6) 다 끝나면 Label.txt 를 train_label.txt(대부분)/val_label.txt(15% 정도)로 나눠 저장",
        "   -> finetune/runpod/det_v5_finetune.yml 이 이 두 파일을 참조함",
        "* 박스 라벨링은 자동화 범위 밖 (사람이 직접). 이 스크립트는 스켈레톤만 만듦",
    ]
    for s in steps:
        ws.append([s])
    ws2 = wb.create_sheet("대상 이미지")
    ws2.append(["image_id", "완료"])
    for c in range(1, 3):
        ws2.cell(1, c).font = Font(bold=True)
    for iid in image_ids:
        ws2.append([iid, ""])
    if missing:
        ws3 = wb.create_sheet("이미지 없음")
        ws3.append(["image_id"])
        for iid in missing:
            ws3.append([iid])
    wb.save(os.path.join(data_dir, "det", "labeling_instructions.xlsx"))


if __name__ == "__main__":
    main()
