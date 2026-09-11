# -*- coding: utf-8 -*-
"""골드셋에서 rec(문자 인식) 파인튜닝용 학습 데이터를 뽑는다.

python finetune/prep_rec_data.py [--limit N] [--out finetune/data]

각 골드 이미지에 대해 v4(960px), v5(960px) det 로 크롭을 얻고(예측이 틀렸거나
stage 가 안 좋은 "hard" 이미지는 1920px+CLAHE v4-loose det 도 추가), 크롭마다
ko/ch rec 을 모두 돌려 dateparse 로 골드 날짜와 일치하는지 본다.

- 일치 + (ko/ch 텍스트 완전 일치 또는 숫자/구분자 비율 70% 이상) -> auto 양성
  (rec/auto/<id>_<k>.png, 라벨 = 일치를 만든 rec 텍스트 그대로)
- 일치했지만 텍스트가 불확실 -> hard (수동 확인 필요, 이미 날짜는 맞지만 라벨 문자열이 애매)
- hard 이미지에서 날짜로 안 풀린 크롭 중 숫자 3자리 이상(또는 fail 이미지의 v5/1920 크롭 전부)
  -> hard (수동 라벨링 대상)

출력은 전부 finetune/data/ 아래. 원본 이미지나 크롭은 git에 안 올라간다 (.gitignore).
"""
from __future__ import annotations

import argparse
import csv
import os
import random
import sys

import cv2

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.dirname(_HERE)
sys.path.insert(0, os.path.join(_REPO, "src"))

import dateparse  # noqa: E402
import ocr  # noqa: E402

_IMAGES_DIR = os.path.normpath(os.path.join(_REPO, os.pardir, "images"))
_GOLD_CSV = os.path.join(_REPO, "labels", "gold.csv")
_PRED_CSV = os.path.join(_REPO, "labels", "auto_gold_v4.csv")
_EXTS = (".jpg", ".jpeg", ".png")
_BAD_STAGES = {"fail", "clahe", "hires", "up2x", "det5"}
_SEP_CHARS = set(".-/: ,")


# --- 유틸 ---------------------------------------------------------------

def _index_images(images_dir: str) -> dict:
    idx = {}
    for f in os.listdir(images_dir):
        stem, ext = os.path.splitext(f)
        if ext.lower() in _EXTS:
            idx[stem] = os.path.join(images_dir, f)
    return idx


def _cand_str(c: dateparse.Candidate) -> tuple:
    y = "NONE" if c.y is None else str(c.y).zfill(4)
    m = "NONE" if c.m is None else str(c.m).zfill(2)
    d = "NONE" if c.d is None else str(c.d).zfill(2)
    return y, m, d


def _matches_gold(c: dateparse.Candidate, gy: str, gm: str, gd: str) -> bool:
    return _cand_str(c) == (gy, gm, gd)


def _digit_sep_ratio(text: str) -> float:
    if not text:
        return 0.0
    n = sum(1 for ch in text if ch.isdigit() or ch in _SEP_CHARS)
    return n / len(text)


def _save_png(img, path: str) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    ok, buf = cv2.imencode(".png", img)
    if not ok:
        raise RuntimeError(f"encode failed: {path}")
    buf.tofile(path)


def _is_wrong(pred_row: dict | None, gy: str, gm: str, gd: str) -> bool:
    if pred_row is None:
        return True
    return (pred_row.get("year") != gy or pred_row.get("month") != gm
            or pred_row.get("day") != gd)


# --- 크롭 수집 -----------------------------------------------------------

def _collect_items(path: str, run_extra: bool) -> list:
    """[(crop, box, source)]. source in {v4, v5, hi}."""
    img = ocr._load(path)
    items = [(c, b, "v4") for c, b in ocr._crops(img)]
    items += [(c, b, "v5") for c, b in ocr._crops(img, v5=True)]
    if run_extra:
        hi_img = ocr._clahe(ocr._load(path, 1920))
        items += [(c, b, "hi") for c, b in ocr._crops(hi_img, loose=True)]
    return items


def _rec_texts(crops: list) -> tuple:
    if not crops:
        return [], []
    ko = [t for t, _ in ocr._rec("ko")(crops)[0]]
    ch = [t for t, _ in ocr._rec("ch")(crops)[0]]
    return ko, ch


# --- 이미지 1장 처리 -------------------------------------------------------

def process_image(image_id: str, gold_row: dict, pred_row: dict | None,
                   img_path: str, out_dir: str) -> dict:
    gy, gm, gd = gold_row["year"], gold_row["month"], gold_row["day"]
    stage = (pred_row or {}).get("stage", "fail")
    pred_hard = _is_wrong(pred_row, gy, gm, gd) or stage in _BAD_STAGES

    items = _collect_items(img_path, run_extra=pred_hard)
    crops = [c for c, _, _ in items]
    ko_texts, ch_texts = _rec_texts(crops)

    auto_rows: list = []   # (relpath, label)
    hard_rows: list = []   # dict rows for xlsx
    matched_any = False
    k = 0
    for (crop, _box, src), ko_t, ch_t in zip(items, ko_texts, ch_texts):
        m_ko = any(_matches_gold(c, gy, gm, gd) for c in dateparse.find_candidates(ko_t))
        m_ch = any(_matches_gold(c, gy, gm, gd) for c in dateparse.find_candidates(ch_t))
        if m_ko or m_ch:
            matched_any = True
            label_text = ko_t if m_ko else ch_t
            agree = ko_t.strip() == ch_t.strip()
            conf_ok = agree or _digit_sep_ratio(label_text) >= 0.7
            fname = f"{image_id}_{k}.png"
            if conf_ok:
                _save_png(crop, os.path.join(out_dir, "rec", "auto", fname))
                auto_rows.append((f"rec/auto/{fname}", label_text, image_id))
            else:
                _save_png(crop, os.path.join(out_dir, "rec", "hard", fname))
                hard_rows.append(dict(file=fname, image_id=image_id,
                                       gold_date=f"{gy}-{gm}-{gd}", ko_text=ko_t, ch_text=ch_t,
                                       note="날짜는 맞았지만 ko/ch 불일치+숫자비율<70%: 정확한 표기 확인 필요"))
            k += 1
        elif pred_hard:
            digits = max(sum(ch.isdigit() for ch in ko_t), sum(ch.isdigit() for ch in ch_t))
            force = stage == "fail" and src in ("v5", "hi")
            if digits >= 3 or force:
                fname = f"{image_id}_{k}.png"
                _save_png(crop, os.path.join(out_dir, "rec", "hard", fname))
                hard_rows.append(dict(file=fname, image_id=image_id,
                                       gold_date=f"{gy}-{gm}-{gd}", ko_text=ko_t, ch_text=ch_t,
                                       note=""))
                k += 1

    is_hard_image = pred_hard or not matched_any
    return dict(auto=auto_rows, hard=hard_rows, is_hard_image=is_hard_image)


# --- xlsx 출력 ------------------------------------------------------------

def _write_hard_xlsx(hard_rows: list, out_dir: str) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "hard"
    header = ["file", "image_id", "gold_date", "ko_text", "ch_text", "label", "skip", "note", "미리보기"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).fill = PatternFill("solid", fgColor="DDDDDD")
    widths = [22, 10, 12, 26, 26, 20, 6, 40, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    hard_dir = os.path.join(out_dir, "rec", "hard")
    for i, r in enumerate(hard_rows, start=2):
        vals = [r["file"], r["image_id"], r["gold_date"], r["ko_text"], r["ch_text"], "", "", r["note"]]
        for c, v in enumerate(vals, start=1):
            ws.cell(i, c, v)
        png_path = os.path.join(hard_dir, r["file"])
        try:
            img = XLImage(png_path)
            h_px = img.height
            ws.row_dimensions[i].height = max(20, h_px * 0.75 + 6)
            img.anchor = f"I{i}"
            ws.add_image(img)
        except Exception as e:  # 손상된 png 등, 라벨링엔 지장 없으므로 건너뜀
            ws.cell(i, 9, f"(미리보기 실패: {e})")
    wb.save(os.path.join(out_dir, "rec", "hard_labels.xlsx"))


# --- 메인 ------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None, help="골드 앞에서 N개만")
    ap.add_argument("--out", default=os.path.join(_HERE, "data"))
    args = ap.parse_args()
    out_dir = os.path.abspath(args.out)
    os.makedirs(out_dir, exist_ok=True)

    gold_rows = list(csv.DictReader(open(_GOLD_CSV, encoding="utf-8")))
    if args.limit:
        gold_rows = gold_rows[:args.limit]
    gold_rows = [r for r in gold_rows if (r["year"], r["month"], r["day"]) != ("NONE", "NONE", "NONE")]

    pred_by_id = {r["image_id"]: r for r in csv.DictReader(open(_PRED_CSV, encoding="utf-8"))}
    img_index = _index_images(_IMAGES_DIR)

    all_auto: list = []
    all_hard: list = []
    hard_images: list = []
    covered = 0
    missing = 0
    t0 = __import__("time").time()
    for i, row in enumerate(gold_rows, 1):
        image_id = row["image_id"]
        path = img_index.get(image_id)
        if path is None:
            missing += 1
            continue
        res = process_image(image_id, row, pred_by_id.get(image_id), path, out_dir)
        all_auto.extend(res["auto"])
        all_hard.extend(res["hard"])
        if res["is_hard_image"]:
            hard_images.append(image_id)
        covered += 1
        if i % 20 == 0 or i == len(gold_rows):
            dt = __import__("time").time() - t0
            print(f"{i}/{len(gold_rows)}  {dt/i:.2f}s/img  auto={len(all_auto)} hard={len(all_hard)}", flush=True)

    # auto_label.txt (전체 auto 양성)
    with open(os.path.join(out_dir, "rec", "auto_label.txt"), "w", encoding="utf-8") as f:
        for relpath, label, _iid in all_auto:
            f.write(f"{relpath}\t{label}\n")

    # train/val 분할: image_id 기준 15% 홀드아웃 (seed 42), 이미지 겹침 방지
    ids = sorted({iid for _, _, iid in all_auto})
    rng = random.Random(42)
    rng.shuffle(ids)
    n_val = max(1, round(len(ids) * 0.15)) if ids else 0
    val_ids = set(ids[:n_val])
    train_pairs = [(p, l) for p, l, iid in all_auto if iid not in val_ids]
    val_pairs = [(p, l) for p, l, iid in all_auto if iid in val_ids]
    with open(os.path.join(out_dir, "rec", "train_label.txt"), "w", encoding="utf-8") as f:
        for p, l in train_pairs:
            f.write(f"{p}\t{l}\n")
    with open(os.path.join(out_dir, "rec", "val_label.txt"), "w", encoding="utf-8") as f:
        for p, l in val_pairs:
            f.write(f"{p}\t{l}\n")

    # hard_labels.xlsx + hard_images.txt
    if all_hard:
        _write_hard_xlsx(all_hard, out_dir)
    with open(os.path.join(out_dir, "rec", "hard_images.txt"), "w", encoding="utf-8") as f:
        for iid in sorted(set(hard_images)):
            f.write(iid + "\n")

    print(f"\n===== 완료 =====")
    print(f"골드 대상: {len(gold_rows)}  이미지 커버: {covered}  이미지 없음: {missing}")
    print(f"auto 양성: {len(all_auto)}  (train {len(train_pairs)} / val {len(val_pairs)}, val 이미지 {len(val_ids)}개)")
    print(f"hard 크롭: {len(all_hard)}  hard 이미지: {len(set(hard_images))}")
    print(f"출력: {out_dir}")


if __name__ == "__main__":
    main()
