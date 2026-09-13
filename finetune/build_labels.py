# -*- coding: utf-8 -*-
"""수동 라벨 엑셀에서 사용자가 채운 라벨을 train_label.txt 에 합친다.

python finetune/build_labels.py [--data finetune/data] [--val-frac 0.0]
python finetune/build_labels.py --data finetune/data_v2

- data_v2/rec/manual_labels.xlsx (정답 컬럼, rec/manual/<file>) 가 있으면 그걸 쓰고,
  없으면 data/rec/hard_labels.xlsx (label 컬럼, rec/hard/<file>) 로 예전처럼 동작한다.
- 정답/label 이 비어있거나 "제외"거나, skip 컬럼이 있고 x/o/1/true 면 건너뜀
- 라벨 문자가 korean_PP-OCRv5_rec_mobile.onnx 의 character 메타데이터(ONNX 딕셔너리)에
  없는 문자를 포함하면 경고하고 건너뜀 (공백은 PP-OCRv5 계열이 use_space_char=True 라 허용)
- 통과한 라벨은 <crop_dir>/<file>\t<label> 로 train_label.txt 에 append (--val-frac > 0 이면
  이미지 단위로 일부를 val_label.txt 로 돌림, seed 42)
"""
from __future__ import annotations

import argparse
import os
import random

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.dirname(_HERE)
_REC_ONNX = os.path.join(_REPO, "weights", "korean_PP-OCRv5_rec_mobile.onnx")
_SKIP_VALUES = {"x", "o", "1", "true", "yes", "skip"}
_EXCLUDE_VALUES = {"제외", "exclude"}


def _load_dict_chars(onnx_path: str) -> set:
    import onnxruntime as ort
    sess = ort.InferenceSession(onnx_path, providers=["CPUExecutionProvider"])
    meta = sess.get_modelmeta().custom_metadata_map
    chars = set("".join(meta["character"].split("\n")))
    chars.add(" ")  # PP-OCRv5 계열은 학습 시 use_space_char=True 로 공백을 별도로 덧붙인다
    return chars


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=os.path.join(_HERE, "data"))
    ap.add_argument("--val-frac", type=float, default=0.0)
    args = ap.parse_args()
    data_dir = os.path.abspath(args.data)
    manual_path = os.path.join(data_dir, "rec", "manual_labels.xlsx")
    hard_path = os.path.join(data_dir, "rec", "hard_labels.xlsx")
    if os.path.exists(manual_path):
        xlsx_path, label_col, id_col, crop_subdir = manual_path, "정답", "이미지 번호", "rec/manual"
    elif os.path.exists(hard_path):
        xlsx_path, label_col, id_col, crop_subdir = hard_path, "label", "image_id", "rec/hard"
    else:
        print(f"없음: {manual_path} 또는 {hard_path} (prep_rec_data*.py 먼저 실행)")
        return

    allowed = None
    if os.path.exists(_REC_ONNX):
        allowed = _load_dict_chars(_REC_ONNX)
    else:
        print(f"경고: {_REC_ONNX} 없음. 문자 검증을 건너뜀 (weights/ 를 download_weights.sh 로 받으세요)")

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    header = [c.value for c in ws[1]]
    idx = {name: i for i, name in enumerate(header)}

    kept, skipped_empty, skipped_flag, skipped_bad_char = [], 0, 0, 0
    for r in rows:
        if r is None or r[idx["file"]] is None:
            continue
        file_ = str(r[idx["file"]]).strip()
        image_id = str(r[idx[id_col]]).strip() if r[idx[id_col]] is not None else ""
        label = r[idx[label_col]]
        if "skip" in idx:
            skip = r[idx["skip"]]
            if skip is not None and str(skip).strip().lower() in _SKIP_VALUES:
                skipped_flag += 1
                continue
        if label is None or not str(label).strip() or str(label).strip().lower() in _EXCLUDE_VALUES:
            skipped_empty += 1
            continue
        label = str(label).strip()
        if allowed is not None and any(ch not in allowed for ch in label):
            bad = sorted({ch for ch in label if ch not in allowed})
            print(f"경고: {file_} 라벨 '{label}' 에 딕셔너리에 없는 문자 {bad} -> 건너뜀")
            skipped_bad_char += 1
            continue
        kept.append((file_, label, image_id))

    train_path = os.path.join(data_dir, "rec", "train_label.txt")
    val_path = os.path.join(data_dir, "rec", "val_label.txt")

    if args.val_frac > 0 and kept:
        ids = sorted({iid for _, _, iid in kept})
        rng = random.Random(42)
        rng.shuffle(ids)
        n_val = max(1, round(len(ids) * args.val_frac))
        val_ids = set(ids[:n_val])
    else:
        val_ids = set()

    n_train = n_val_written = 0
    with open(train_path, "a", encoding="utf-8") as ft, open(val_path, "a", encoding="utf-8") as fv:
        for file_, label, iid in kept:
            line = f"{crop_subdir}/{file_}\t{label}\n"
            if iid in val_ids:
                fv.write(line)
                n_val_written += 1
            else:
                ft.write(line)
                n_train += 1

    print(f"라벨링된 행: {len(kept)}  (train {n_train} / val {n_val_written})")
    print(f"건너뜀: 라벨 없음 {skipped_empty}, skip 표시 {skipped_flag}, 딕셔너리 밖 문자 {skipped_bad_char}")


if __name__ == "__main__":
    main()
